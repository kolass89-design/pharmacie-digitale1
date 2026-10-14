#!/usr/bin/env python3
"""MYP 5 Mathematics — Mid-Term Test 2 (Criteria A, B, D) — Oak International School.

Outputs:
  MYP5_Mathematics_MidTerm_Test2.pdf            (print-ready, exactly 1 x A4 portrait)
  MYP5_Mathematics_MidTerm_Test2.docx           (editable)
Requires assets/ois-logo-trimmed.png. All mathematics pre-verified.
"""
import os, math
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor, white
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "ois-logo-trimmed.png")
STRIP = os.path.join(HERE, "assets", "beds_strip.png")

GREEN      = HexColor("#008968")
GREEN_DARK = HexColor("#006B52")
INK        = HexColor("#1A1A1A")
DARK       = HexColor("#1F2A24")
GREY       = HexColor("#6B7280")
LIGHTLINE  = HexColor("#9CA3AF")
DOTCOL     = HexColor("#B3BAC2")
TINT       = HexColor("#E6F4EE")
BARTEXT    = white

W, H = A4
ML = 34; MR = 34
CW = W - ML - MR
XR = W - MR

BODY, LEAD = 9.4, 11.0
DOTSP = 10.2

def rich_width(parts, size):
    w = 0.0
    for text, st in parts:
        f, s = ("Helvetica-Bold", size) if st == "b" else \
               ("Helvetica-Oblique", size) if st == "i" else \
               ("Helvetica", size * 0.68) if st == "sub" else ("Helvetica", size)
        w += stringWidth(text, f, s)
    return w

def draw_rich(c, x, y, parts, size=BODY, color=INK):
    c.setFillColor(color)
    cur = x
    for text, st in parts:
        y_use = y
        if st == "b":
            c.setFont("Helvetica-Bold", size)
        elif st == "i":
            c.setFont("Helvetica-Oblique", size)
        elif st == "sub":
            c.setFont("Helvetica", size * 0.66); y_use = y - 2.4
        else:
            c.setFont("Helvetica", size)
        c.drawString(cur, y_use, text)
        cur += stringWidth(text, c._fontname, c._fontsize)
    return cur

def wrap_parts(parts, size, maxw):
    lines, cur, curw = [], [], 0.0
    for text, st in parts:
        words = text.split(" ")
        for i, wd in enumerate(words):
            piece = wd + (" " if i < len(words) - 1 else "")
            pw = stringWidth(piece, "Helvetica-Bold" if st == "b" else "Helvetica", size)
            if cur and curw + pw > maxw:
                lines.append(cur); cur, curw = [], 0.0
            cur.append((piece, st)); curw += pw
    if cur: lines.append(cur)
    return lines

# ---------------------------------------------------------------- figures strip
def make_beds_strip(path):
    """Triangular stone-edged beds: Bed 1 (3 stones), Bed 2 (6), Bed 3 (9)."""
    SC = 8                       # px per pt
    sp = 11.5                    # dot spacing in pt
    r = 2.7                      # dot radius pt
    gap = 30.0                   # gap between beds pt
    lab_h = 13.0                 # label band pt
    pad = 3.0
    beds = [2, 3, 4]             # dots per side for beds 1..3
    widths, dots = [], []
    for s in beds:
        w = (s - 1) * sp
        h = math.sqrt(3) / 2 * w
        pts = set()
        import itertools
        base_l, base_r, apex = (0.0, h), (w, h), (w / 2, 0.0)
        for k in range(s):
            t = k / (s - 1)
            pts.add((round(base_l[0] + t * (apex[0] - base_l[0]), 2),
                     round(base_l[1] + t * (apex[1] - base_l[1]), 2)))
            pts.add((round(apex[0] + t * (base_r[0] - apex[0]), 2),
                     round(apex[1] + t * (base_r[1] - apex[1]), 2)))
            pts.add((round(base_l[0] + t * (base_r[0] - base_l[0]), 2),
                     round(base_l[1] + t * (base_r[1] - base_l[1]), 2)))
        widths.append((w, h)); dots.append(sorted(pts))
    total_w = sum(w for w, _ in widths) + gap * 2 + pad * 2
    max_h = max(h for _, h in widths)
    total_h = pad + max_h + r + 2.5 + lab_h + pad  # common label baseline under everything
    img = Image.new("RGB", (int(total_w * SC), int(total_h * SC)), "white")
    d = ImageDraw.Draw(img)
    try:
        fnt = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", int(6.6 * SC))
    except OSError:
        fnt = ImageFont.load_default()
    label_y = pad + max_h + r + 2.5
    xoff = pad
    for i, (w, h) in enumerate(widths):
        cx = xoff + w / 2
        yshift = (max_h - h) / 2  # vertically centre smaller triangles
        for (px, py) in dots[i]:
            X, Y = (xoff + px), (pad + yshift + py)
            d.ellipse([ (X - r) * SC, (Y - r) * SC, (X + r) * SC, (Y + r) * SC ],
                      fill=(0, 137, 104), outline=(0, 96, 74), width=int(0.9 * SC))
        lab = f"Bed {i + 1}"
        tw = d.textlength(lab, font=fnt)
        d.text((cx * SC - tw / 2, label_y * SC), lab, fill=(31, 42, 36), font=fnt)
        xoff += w + gap
    img.save(path)
    return total_w, total_h

STRIP_W, STRIP_H = make_beds_strip(STRIP)

# ---------------------------------------------------------------- paper class
class Paper:
    def __init__(self, path):
        self.c = c = canvas.Canvas(path, pagesize=A4)
        c.setTitle("MYP 5 Mathematics - Mid-Term Test 2")
        c.setAuthor("Oak International School")
        c.setSubject("Linear and Arithmetic Sequences - Criteria A, B and D")
        self.y = H - ML

    def header(self):
        c = self.c
        logo_h = 42.0
        logo_w = logo_h * 0.787
        c.drawImage(LOGO, ML, self.y - logo_h, width=logo_w, height=logo_h,
                    mask="auto", preserveAspectRatio=True)
        tx = ML + logo_w + 10
        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 14.5)
        c.drawString(tx, self.y - 12, "OAK INTERNATIONAL SCHOOL", charSpace=1.1)
        c.setFillColor(DARK)
        c.setFont("Helvetica-Bold", 10.3)
        c.drawString(tx, self.y - 26, "MYP 5 MATHEMATICS — MID-TERM TEST 2", charSpace=0.3)
        c.setFillColor(GREY)
        c.setFont("Helvetica-Oblique", 7.2)
        c.drawString(tx, self.y - 37, "Bilingual Excellence in Education & Leadership")
        bw, bh = 132, 41
        bx, by = XR - bw, self.y - bh - 1
        c.setStrokeColor(GREEN); c.setLineWidth(0.9)
        c.roundRect(bx, by, bw, bh, 2.5, stroke=1, fill=0)
        rows = [("DURATION", "30 minutes"), ("TOTAL MARKS", "20"),
                ("CALCULATOR", "Not permitted")]
        yy = by + bh - 11.5
        for lab, val in rows:
            c.setFont("Helvetica-Bold", 6.4); c.setFillColor(GREY)
            c.drawString(bx + 7, yy, lab)
            c.setFont("Helvetica-Bold", 7.5); c.setFillColor(DARK)
            c.drawRightString(bx + bw - 7, yy, val)
            yy -= 10.2
        yl = by - 13
        c.setFont("Helvetica-Bold", 8.4); c.setFillColor(DARK)
        segs = [("Name:", 30, 205), ("Class:", 250, 385), ("Date:", 430, XR)]
        for lab, x0, x1 in segs:
            c.drawString(x0, yl, lab)
            lw = stringWidth(lab, "Helvetica-Bold", 8.4)
            c.setStrokeColor(LIGHTLINE); c.setLineWidth(0.6)
            c.line(x0 + lw + 4, yl - 1.5, x1, yl - 1.5)
        self.y = yl - 8
        c.setStrokeColor(GREEN); c.setLineWidth(1.3)
        c.line(ML, self.y, XR, self.y)
        self.y -= 9

    def instructions(self):
        c = self.c
        lines = [
            "Answer ALL questions. Show all working — marks are awarded for method, for general rules, for verification of rules, and for interpretation in context.",
            "Write final answers on the dotted lines. Give non-exact decimal answers correct to 2 d.p. Amounts in francs are exact. Aim for about 90 seconds per mark.",
        ]
        h = 5.5 + LEAD * len(lines) + 3.5
        c.setFillColor(TINT)
        c.roundRect(ML, self.y - h, CW, h, 2, stroke=0, fill=1)
        c.setFillColor(GREEN)
        c.rect(ML, self.y - h, 2.6, h, stroke=0, fill=1)
        yy = self.y - 5.5 - 8
        c.setFillColor(DARK); c.setFont("Helvetica", 7.35)
        for ln in lines:
            c.drawString(ML + 10, yy, ln)
            yy -= 9.6
        self.y -= h + 9.5

    def task_bar(self, title, crit, marks):
        c = self.c
        bh = 13.6
        c.setFillColor(GREEN)
        c.roundRect(ML, self.y - bh, CW, bh, 2, stroke=0, fill=1)
        c.setFillColor(BARTEXT)
        c.setFont("Helvetica-Bold", 7.7)
        c.drawString(ML + 8, self.y - bh + 4.2, title, charSpace=0.2)
        right = f"{crit}  •  {marks} MARKS"
        c.setFont("Helvetica-Bold", 7.2)
        c.drawRightString(XR - 8, self.y - bh + 4.3, right)
        self.y -= bh + 4.6

    def para(self, parts, size=BODY, color=INK, indent=0, hang=0):
        lines = wrap_parts(parts, size, CW - indent - hang)
        for i, ln in enumerate(lines):
            self.y -= LEAD
            draw_rich(self.c, ML + indent + (hang if i > 0 else 0), self.y, ln, size, color)
        return lines

    def seq_display(self, groups):
        c = self.c
        parts = []
        for gi, (lab, terms) in enumerate(groups):
            if gi: parts.append(("     ", ""))
            parts.append((lab + ": ", "b"))
            parts.append((", ".join(terms) + ",   …", ""))
        wd = rich_width(parts, BODY + 0.4)
        self.y -= LEAD + 1.4
        draw_rich(c, ML + (CW - wd) / 2, self.y, parts, BODY + 0.4)
        self.y -= 1.5

    def picture(self, path, w_pt, h_pt):
        self.y -= h_pt + 2
        self.c.drawImage(path, ML + (CW - w_pt) / 2, self.y, width=w_pt, height=h_pt,
                         mask="auto")
        self.y -= 3.5

    def subq(self, label, parts, marks, dots=1):
        c = self.c
        parts = [(f"({label})  ", "b")] + parts
        lw = stringWidth(f"({label})  ", "Helvetica-Bold", BODY)
        self.para(parts, hang=lw)
        for i in range(dots):
            self.y -= DOTSP
            c.setStrokeColor(DOTCOL); c.setLineWidth(0.7)
            c.setDash(0.6, 2.6)
            c.line(ML + lw, self.y, XR - 52, self.y)
            c.setDash()
        c.setFont("Helvetica-Oblique", 7.6); c.setFillColor(GREY)
        c.drawRightString(XR, self.y + 1.2, f"[{marks} mark{'' if marks == 1 else 's'}]")
        self.y -= 3.2

    def footer(self):
        c = self.c
        c.setStrokeColor(HexColor("#D5DBE0")); c.setLineWidth(0.6)
        c.line(ML, 40, XR, 40)
        c.setFont("Helvetica", 6.7); c.setFillColor(GREY)
        c.drawString(ML, 30.5, "Oak International School — MYP 5 Mathematics — Mid-Term Test 2")
        c.drawRightString(XR, 30.5, "Page 1 of 1")

    def space(self, pts):
        self.y -= pts

    def save(self):
        self.c.save()


p = Paper(os.path.join(HERE, "MYP5_Mathematics_MidTerm_Test2.pdf"))
p.header()
p.instructions()

# ---------------------------------------------------------------- TASK 1 (A, 4)
p.task_bar("TASK 1 — UNDERSTANDING NUMBER PATTERNS", "Criterion A", 4)
p.para([("For a class display, three students each recorded a number pattern:", "")])
p.seq_display([("Ali", ["4", "10", "16", "22"]), ("Bea", ["1", "4", "9", "16"]),
               ("Camille", ["40", "36", "32", "28"])])
p.subq("a", [("Two of the three patterns change by a constant amount. Identify these two patterns and state the constant change in each.", "")], 2, dots=1)
p.subq("b", [("A classmate says that Bea’s pattern is linear ", ""),
             ("“because it always increases.”", "i"),
             (" Explain the classmate’s error, using calculations.", "")], 2, dots=2)
p.space(7)

# ---------------------------------------------------------------- TASK 2 (B, 6)
p.task_bar("TASK 2 — INVESTIGATION: STONE-EDGED GARDEN BEDS", "Criterion B", 6)
p.para([("For the school’s greening project, the mathematics club builds triangular garden beds edged with equal stones. Beds 1, 2 and 3 are shown, in order.", "")])
p.picture(STRIP, STRIP_W, STRIP_H)
p.subq("a", [("Determine the number of stones needed for Bed 4 and Bed 5.", "")], 2, dots=1)
p.subq("b", [("The club describes the pattern with a general rule. Write a formula for the number of stones, ", ""),
             ("S", "i"), (", in Bed ", ""), ("n", "i"), (", where ", ""), ("n", "i"),
             (" is the bed number.", "")], 2, dots=1)
p.subq("c", [("The gardener has ordered 30 stones — according to the club, exactly enough for Bed 10. Use your rule to verify the order, and explain why your rule must work for ", ""),
             ("every", "i"),
             (" bed, referring to the structure of the pattern rather than only to the numbers.", "")], 2, dots=2)
p.space(7)

# ---------------------------------------------------------------- TASK 3 (D, 6)
p.task_bar("TASK 3 — APPLICATION: AMINA’S SAVINGS CHALLENGE", "Criterion D", 6)
p.para([("Amina joins a one-year savings challenge in Cotonou. She deposits 10 000 francs in January (month 1), and from then on she increases each month’s deposit by exactly 2 500 francs.", "")])
p.subq("a", [("Show that Amina deposits (2 500", ""), ("n", "i"),
             (" + 7 500) francs in month ", ""), ("n", "i"), (".", "")], 2, dots=1)
p.subq("b", [("How much does Amina deposit in December (month 12)?", "")], 1, dots=1)
p.subq("c", [("Calculate the total amount Amina has saved by the end of December.", "")], 2, dots=2)
p.subq("d", [("In January, Amina sees a bicycle costing 300 000 francs, payable at the end of December. Determine whether her savings will be enough. Justify with a calculation.", "")], 1, dots=1)
p.space(7)

# ---------------------------------------------------------------- TASK 4 (A, 4)
p.task_bar("TASK 4 — REASONING CHALLENGE: TWO CANDLES", "Criterion A", 4)
p.para([("Two candles are lit at the same moment. Candle A is 24 cm tall and burns 1.5 cm every hour; Candle B is 16 cm tall and burns 1 cm every hour. A student claims: ", ""),
        ("“Candle B will still be burning after Candle A has gone out.”", "i")])
p.subq("a", [("Write a formula for the height of each candle after ", ""), ("t", "i"), (" hours.", "")], 1, dots=1)
p.subq("b", [("Investigate the student’s claim. Find when the two candles have the same height, and state what your result means for the claim. Justify fully.", "")], 3, dots=3)

p.footer()
p.save()
print("PDF 2 written")

# ================================================================== DOCX
GREEN_HEX = "008968"

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_shading(par, hexfill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexfill)
    par._p.get_or_add_pPr().append(shd)

def bottom_border(par, color=GREEN_HEX, size=12):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(size))
    b.set(qn('w:space'), '2'); b.set(qn('w:color'), color)
    pbdr.append(b); pPr.append(pbdr)

def dotted_line(par, pos_twips=9900):
    pPr = par._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right'); tab.set(qn('w:leader'), 'dot'); tab.set(qn('w:pos'), str(pos_twips))
    tabs.append(tab); pPr.append(tabs)
    r = par.add_run('\t'); r.font.size = Pt(9)

def add_rich(par, parts, size=10):
    for text, st in parts:
        r = par.add_run(text)
        r.font.size = Pt(size * (0.66 if st == "sub" else 1.0))
        r.font.bold = (st == "b"); r.font.italic = (st == "i")
        if st == "sub": r.font.subscript = True
    return par

def styled_par(doc, space_after=2, space_before=0, align=None, line=1.0):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if align: par.alignment = align
    return par

def build_student_docx(path):
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = Cm(1.15); sec.bottom_margin = Cm(1.0)
        sec.left_margin = Cm(1.35); sec.right_margin = Cm(1.35)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(2)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Cm(1.9); tbl.columns[1].width = Cm(16.4)
    c0 = tbl.cell(0, 0).paragraphs[0]
    c0.add_run().add_picture(LOGO, height=Cm(1.5))
    c1 = tbl.cell(0, 1)
    p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run("OAK INTERNATIONAL SCHOOL")
    r.font.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor(0x00, 0x89, 0x68)
    p2 = c1.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run("MYP 5 MATHEMATICS — MID-TERM TEST 2")
    r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x24)
    p3 = c1.add_paragraph(); p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run("Bilingual Excellence in Education & Leadership")
    r.font.italic = True; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    meta = styled_par(doc, space_before=3, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    r = meta.add_run("Duration: 30 minutes   |   Total: 20 marks   |   Calculator: not permitted")
    r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = RGBColor(0x00, 0x6B, 0x52)

    nd = styled_par(doc, space_after=4)
    r = nd.add_run("Name: _____________________________        Class: ________________        Date: ______________")
    r.font.size = Pt(9)

    rule = styled_par(doc, space_after=5); bottom_border(rule)

    ins = styled_par(doc, space_after=2)
    set_shading(ins, "E6F4EE")
    r = ins.add_run("Answer ALL questions. Show all working — marks are awarded for method, for general rules, for verification of rules, and for interpretation in context. "
                    "Write final answers on the dotted lines. Non-exact decimals: 2 d.p. Amounts in francs are exact. Aim for about 90 seconds per mark.")
    r.font.size = Pt(7.6)

    def task(title, crit, marks):
        par = styled_par(doc, space_before=5, space_after=3)
        set_shading(par, GREEN_HEX)
        r = par.add_run(f"  {title}")
        r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r2 = par.add_run(f"   —   {crit} • {marks} MARKS  ")
        r2.font.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(0xD9, 0xF2, 0xE9)

    def subq(label, parts, marks, dots=1):
        par = styled_par(doc)
        add_rich(par, [(f"({label})  ", "b")] + parts)
        mr = par.add_run(f"\t[{marks} mark{'' if marks == 1 else 's'}]")
        mr.font.size = Pt(8); mr.font.italic = True; mr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        pPr = par._p.get_or_add_pPr()
        tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), '10100')
        tabs.append(tab); pPr.append(tabs)
        for _ in range(dots):
            d = styled_par(doc, space_after=1)
            d.paragraph_format.left_indent = Cm(0.45)
            dotted_line(d)

    task("TASK 1 — UNDERSTANDING NUMBER PATTERNS", "Criterion A", 4)
    par = styled_par(doc); par.add_run("For a class display, three students each recorded a number pattern:")
    sq = styled_par(doc, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = sq.add_run("Ali:  4, 10, 16, 22, …     Bea:  1, 4, 9, 16, …     Camille:  40, 36, 32, 28, …")
    r.font.size = Pt(10)
    subq("a", [("Two of the three patterns change by a constant amount. Identify these two patterns and state the constant change in each.", "")], 2, dots=1)
    subq("b", [("A classmate says that Bea’s pattern is linear ", ""), ("“because it always increases.”", "i"),
               (" Explain the classmate’s error, using calculations.", "")], 2, dots=2)

    task("TASK 2 — INVESTIGATION: STONE-EDGED GARDEN BEDS", "Criterion B", 6)
    par = styled_par(doc)
    par.add_run("For the school’s greening project, the mathematics club builds triangular garden beds edged with equal stones. Beds 1, 2 and 3 are shown, in order.")
    fig = styled_par(doc, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    fig.add_run().add_picture(STRIP, width=Cm(4.3))
    subq("a", [("Determine the number of stones needed for Bed 4 and Bed 5.", "")], 2, dots=1)
    subq("b", [("The club describes the pattern with a general rule. Write a formula for the number of stones, ", ""), ("S", "i"),
               (", in Bed ", ""), ("n", "i"), (", where ", ""), ("n", "i"), (" is the bed number.", "")], 2, dots=1)
    subq("c", [("The gardener has ordered 30 stones — according to the club, exactly enough for Bed 10. Use your rule to verify the order, and explain why your rule must work for ", ""),
               ("every", "i"),
               (" bed, referring to the structure of the pattern rather than only to the numbers.", "")], 2, dots=2)

    task("TASK 3 — APPLICATION: AMINA’S SAVINGS CHALLENGE", "Criterion D", 6)
    par = styled_par(doc)
    par.add_run("Amina joins a one-year savings challenge in Cotonou. She deposits 10 000 francs in January (month 1), and from then on she increases each month’s deposit by exactly 2 500 francs.")
    subq("a", [("Show that Amina deposits (2 500", ""), ("n", "i"), (" + 7 500) francs in month ", ""), ("n", "i"), (".", "")], 2, dots=1)
    subq("b", [("How much does Amina deposit in December (month 12)?", "")], 1, dots=1)
    subq("c", [("Calculate the total amount Amina has saved by the end of December.", "")], 2, dots=2)
    subq("d", [("In January, Amina sees a bicycle costing 300 000 francs, payable at the end of December. Determine whether her savings will be enough. Justify with a calculation.", "")], 1, dots=1)

    task("TASK 4 — REASONING CHALLENGE: TWO CANDLES", "Criterion A", 4)
    par = styled_par(doc)
    add_rich(par, [("Two candles are lit at the same moment. Candle A is 24 cm tall and burns 1.5 cm every hour; Candle B is 16 cm tall and burns 1 cm every hour. A student claims: ", ""),
                   ("“Candle B will still be burning after Candle A has gone out.”", "i")])
    subq("a", [("Write a formula for the height of each candle after ", ""), ("t", "i"), (" hours.", "")], 1, dots=1)
    subq("b", [("Investigate the student’s claim. Find when the two candles have the same height, and state what your result means for the claim. Justify fully.", "")], 3, dots=3)

    doc.core_properties.title = "MYP 5 Mathematics — Mid-Term Test 2"
    doc.core_properties.author = "Oak International School"
    doc.save(path)
    print("Student DOCX 2 written")

build_student_docx(os.path.join(HERE, "MYP5_Mathematics_MidTerm_Test2.docx"))
