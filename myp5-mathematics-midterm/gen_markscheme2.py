#!/usr/bin/env python3
"""Teacher's Mark Scheme — MYP 5 Mid-Term Test 2 (Criteria A, B, D)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "ois-logo-trimmed.png")

GREEN = RGBColor(0x00, 0x89, 0x68)
DARK  = RGBColor(0x1F, 0x2A, 0x24)
GREY  = RGBColor(0x6B, 0x72, 0x80)
GREEN_HEX = "008968"
TINT_HEX  = "E6F4EE"

doc = Document()
for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(3)

def shade(par, hexfill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexfill)
    par._p.get_or_add_pPr().append(shd)

def shade_cell(cell, hexfill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexfill)
    cell._tc.get_or_add_tcPr().append(shd)

def par(space_after=3, space_before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align: p.alignment = align
    return p

def rich(p, parts, size=10):
    for text, stl in parts:
        r = p.add_run(text)
        r.font.size = Pt(size * (0.66 if stl == 'sub' else 1.0))
        r.font.bold = (stl == 'b'); r.font.italic = (stl == 'i')
        if stl == 'sub': r.font.subscript = True
    return p

# header
tbl = doc.add_table(rows=1, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.columns[0].width = Cm(1.9); tbl.columns[1].width = Cm(16.1)
c0 = tbl.cell(0, 0).paragraphs[0]
c0.add_run().add_picture(LOGO, height=Cm(1.5))
c1 = tbl.cell(0, 1)
p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
r = p.add_run("OAK INTERNATIONAL SCHOOL"); r.font.bold = True; r.font.size = Pt(15); r.font.color.rgb = GREEN
p = c1.add_paragraph(); p.paragraph_format.space_after = Pt(0)
r = p.add_run("MYP 5 MATHEMATICS — MID-TERM TEST 2"); r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = DARK
p = c1.add_paragraph(); p.paragraph_format.space_after = Pt(0)
r = p.add_run("TEACHER'S MARK SCHEME — CONFIDENTIAL"); r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = GREY

p = par(space_before=4, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
r = p.add_run("Duration: 30 minutes   |   Total: 20 marks   |   Criteria: A, B and D")
r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = RGBColor(0x00, 0x6B, 0x52)

def band(text, right=""):
    p = par(space_before=8, space_after=4)
    shade(p, GREEN_HEX)
    r = p.add_run(f"  {text}"); r.font.bold = True; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if right:
        r2 = p.add_run(f"   —   {right}  "); r2.font.bold = True; r2.font.size = Pt(8.5)
        r2.font.color.rgb = RGBColor(0xD9, 0xF2, 0xE9)

def note(text, bold_lead=None):
    p = par(space_after=3)
    if bold_lead:
        r = p.add_run(bold_lead + " "); r.font.bold = True; r.font.size = Pt(8.8); r.font.color.rgb = DARK
    r = p.add_run(text); r.font.size = Pt(8.8); r.font.color.rgb = GREY
    return p

def solution_table(rows):
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(8.6), Cm(1.7), Cm(7.7)]
    hdr = t.rows[0].cells
    for i, htxt in enumerate(["Answer / working", "Marks", "Marking notes"]):
        hdr[i].width = widths[i]
        p = hdr[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(htxt); r.font.bold = True; r.font.size = Pt(8.6)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], GREEN_HEX)
    for working, marks, notes in rows:
        cells = t.add_row().cells
        for i, val in enumerate((working, marks, notes)):
            cells[i].width = widths[i]
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if isinstance(val, str) and val.startswith("RAW:"):
                rich(p, [("  ", "")] + eval(val[4:]), size=9)
            else:
                r = p.add_run(("  " if i != 1 else "") + str(val))
                r.font.size = Pt(9)
                if i == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t

# ================================================================ TASK 1 (A)
band("TASK 1 — UNDERSTANDING NUMBER PATTERNS", "Criterion A • 4 marks")
solution_table([
    ("Ali is linear with constant change +6.\nCamille is linear with constant change −4 (a constant decrease is still linear).\nBea's pattern is the one left out.", "2",
     "A1 for both correct patterns with both correct changes (+6 and −4). A1 for correctly excluding Bea. Accept 'decreases by 4 each time' for Camille. A student who writes 'Camille: −4, so not linear' has the key misconception — award 0 for Camille and flag for feedback."),
    ("The classmate has confused increasing with linear. A linear pattern needs a CONSTANT change between consecutive terms. Bea's differences: 4−1 = 3, 9−4 = 5, 16−9 = 7. Since 3 ≠ 5 ≠ 7, the change is not constant, so Bea's pattern is not linear — even though it increases.", "2",
     "M1 for computing at least two successive differences of Bea's pattern; A1 for the explicit statement that 'increasing' does not imply 'constant change', supported by the differences. (Bea's terms are perfect squares 1², 2², 3², 4² — accept, not required.)"),
])
note("Growth alone never establishes linearity; the constant-difference test is required, for positive and negative differences alike (Task 1a, Camille).",
     "Misconception:")

# ================================================================ TASK 2 (B)
band("TASK 2 — INVESTIGATION: STONE-EDGED GARDEN BEDS", "Criterion B • 6 marks")
solution_table([
    ("Bed 4 = 12 stones, Bed 5 = 15 stones.", "2",
     "A1 for 12; A1 for 15. Accept counting on (+3 each time) or drawing Bed 4. This strand rewards discovering the pattern from the figures (Criterion B i)."),
    ("S = 3n  (equivalently uₙ = 3 + (n − 1) × 3, or 'three times the bed number').", "2",
     "M1 for expressing the rule from the findings (e.g. 'multiply the bed number by 3' or a recursive description +3 starting at 3); A1 for the correct general formula S = 3n. Accept S = 3 + 3(n − 1) unevaluated. (Criterion B ii — describing patterns as general rules.)"),
    ("Verification: S = 3 × 10 = 30 stones, so the order of 30 stones is exactly right. ✓\nExplanation (any one, in student's words):\n• Each new bed adds exactly one stone to each of the three sides, i.e. +3 stones per bed, starting from 3 stones in Bed 1 — so S = 3 + 3(n − 1) = 3n for every n.\n• Bed n has (n + 1) stones along each side; 3 sides × (n + 1) counts each of the 3 corner stones twice, so S = 3(n + 1) − 3 = 3n.", "2",
     "A1 for checking 30 = 3 × 10 and confirming the order; R1 for a structural justification that explains the rule's validity for ALL beds (the +3-per-side growth or the 3 × (n+1) − 3 corner argument). A purely numerical confirmation ('it works for beds 1–4') does not earn R1. (Criterion B iii — verifying validity.)"),
])
note("Level 7–8 in Criterion B requires the general rule AND evidence of verification/justification. Students who only extend the table stop at level 3–4 evidence for strand iii.",
     "Marking guidance:")

# ================================================================ TASK 3 (D)
band("TASK 3 — APPLICATION: AMINA'S SAVINGS CHALLENGE", "Criterion D • 6 marks")
p = par(space_after=3)
rich(p, [("Model: deposits form an arithmetic sequence with ", ""), ("a", "i"), (" = 10 000 (January, month 1) and ", ""), ("d", "i"),
         (" = 2 500 francs per month.", "")], size=9.5)
solution_table([
    ("uₙ = a + (n − 1)d = 10 000 + (n − 1)(2 500)\n= 10 000 + 2 500n − 2 500 = 2 500n + 7 500. ✓\nCheck: month 1 → 2 500 + 7 500 = 10 000 ✓", "2",
     "M1 for translating the situation into the arithmetic model uₙ = a + (n − 1)d with a = 10 000, d = 2 500; A1 for correct simplification to 2 500n + 7 500. Verifying with month 1 (or 2) strengthens the response but is not required."),
    ("December is month 12: u₁₂ = 2 500 × 12 + 7 500 = 30 000 + 7 500 = 37 500 francs.", "1", "A1 for 37 500 francs."),
    ("S₁₂ = 12⁄2 × (u₁ + u₁₂) = 6 × (10 000 + 37 500) = 6 × 47 500 = 285 000 francs.\n(Alternative: S₁₂ = 6 × (2 × 10 000 + 11 × 2 500) = 285 000.)", "2",
     "M1 for a correct total method (sum formula with consistent values, or adding the twelve deposits 10 000 + 12 500 + … + 37 500); A1 for 285 000 francs. '12 × 37 500' assumes the December deposit repeats — 0 marks."),
    ("285 000 < 300 000, so Amina's savings are NOT enough; she is short by 300 000 − 285 000 = 15 000 francs.", "1",
     "A1 for the explicit comparison with a decision and the shortfall (15 000 F). A bare 'no' without the comparison does not answer the contextual question."),
])
note("Criterion D rewards the full cycle: translate (a), apply (b, c), interpret (d) — the comparison and the 15 000 F shortfall are the interpreting step.",
     "Contextual note:")
note("Using n instead of (n − 1) in (a) gives 2 500n + 10 000; taking u₁₂ = 40 000; multiplying the final deposit by 12 in (c).",
     "Likely errors:")

# ================================================================ TASK 4 (A)
band("TASK 4 — REASONING CHALLENGE: TWO CANDLES", "Criterion A • 4 marks")
solution_table([
    ("RAW: [('A:  ', 'b'), ('h', ''), ('A', 'sub'), (' = 24 − 1.5', ''), ('t', 'i'), ('        B:  ', 'b'), ('h', ''), ('B', 'sub'), (' = 16 − ', ''), ('t', 'i'), ('', '')]", "1",
     "A1 for both formulae (any equivalent form, e.g. h = 24 − 3t⁄2). One only: no mark."),
    ("Same height:  24 − 1.5t = 16 − t  →  8 = 0.5t  →  t = 16 hours.\nHeights at t = 16:  A: 24 − 24 = 0 cm;  B: 16 − 16 = 0 cm.\nBurn times: A burns 24 ⁄ 1.5 = 16 hours; B burns 16 ⁄ 1 = 16 hours.\nVerdict: the claim is FALSE. Both candles finish at the same moment (16 hours) — they reach the same height (0 cm) exactly when they both go out, so B is never still burning after A has gone out.", "3",
     "M1 for forming the equation (or comparing burn times 24 ⁄ 1.5 and 16 ⁄ 1, or a table of heights at consistent hourly values); A1 for t = 16 hours with the interpretation height 0 cm — or burn time 16 h for both; R1 for the explicit verdict that the claim is false, justified by the equal finish time (heights at/ before t = 16 may be tabulated as support). Note: the equality occurs only at t = 16, when both are extinguished — answers of '16 hours' without interpreting that both heights are 0 do not earn R1."),
])
note("A student who finds t = 16 and writes 'the candles are equal' without stating that this is the moment both burn out has not interpreted the result — probe with 'what is each candle's height at t = 16?'",
     "Marking guidance:")

# ================================================================ CRITERION MAPPING
band("CRITERION MAPPING")
t = doc.add_table(rows=1, cols=4)
t.style = 'Table Grid'
widths = [Cm(2.2), Cm(1.6), Cm(3.2), Cm(10.0)]
hdr = t.rows[0].cells
for i, htxt in enumerate(["Task", "Marks", "Criterion", "Objective strands addressed"]):
    hdr[i].width = widths[i]
    p = hdr[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(htxt); r.font.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(hdr[i], GREEN_HEX)
mapping = [
    ("Task 1", "4", "A", "i–iii: knowing and applying the constant-change test for linearity; accurate differences; communicating a justified conclusion."),
    ("Task 2", "6", "B", "i–iii: selecting and applying problem-solving techniques to discover the pattern (a); describing the pattern as a general rule (b); verifying the rule's validity and justifying it structurally (c)."),
    ("Task 3", "6", "D", "i–iv: identifying the arithmetic model in a savings context; applying the strategy; accurate solution; interpreting the shortfall in context."),
    ("Task 4", "4", "A", "ii–iii: applying algebraic techniques to model and compare two linear patterns; justifying a verdict."),
    ("Total", "20", "A: 8   B: 6   D: 6", "Criterion A = 40%, Criterion B = 30%, Criterion D = 30%."),
]
for row in mapping:
    cells = t.add_row().cells
    for i, val in enumerate(row):
        cells[i].width = widths[i]
        p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(val); r.font.size = Pt(8.8)
        if row[0] == "Total":
            r.font.bold = True; shade_cell(cells[i], TINT_HEX)

p = par(space_before=4, space_after=2)
r = p.add_run("Note on criterion levels: a single 30-minute test samples the criteria; it contributes evidence towards levels for Criteria A, B and D but should not, alone, determine final achievement levels.")
r.font.italic = True; r.font.size = Pt(8.4); r.font.color.rgb = GREY

# ================================================================ MISCONCEPTIONS
band("COMMON MISCONCEPTIONS AND FEEDBACK NOTES")
misodes = [
    ("'Increasing means linear.'", "Bea's pattern (1, 4, 9, 16) rises yet is not linear; only a constant-difference test decides (Task 1b)."),
    ("'A negative change cannot be constant.'", "Camille's −4 is a perfectly valid common difference; decreasing sequences are linear too (Task 1a)."),
    ("Extending ≠ generalising", "Listing more terms shows the pattern; a rule S = 3n plus a structural argument ('+3 per side, always') is what verifies it (Task 2c)."),
    ("Numerical checks as proof", "Checking a few beds cannot prove 'every bed'; the justification must reference the pattern's structure, not the numbers alone (Task 2c)."),
    ("Total = last term × count", "In Task 3(c), 12 × 37 500 ignores the growth; use the sum formula or add all twelve deposits."),
    ("Stopping at the algebraic answer", "In both contextual tasks, the final mark requires interpretation: the 15 000 F shortfall (T3d) and 'both go out together at 16 h' (T4b)."),
]
t = doc.add_table(rows=1, cols=2)
t.style = 'Table Grid'
widths = [Cm(5.6), Cm(11.4)]
hdr = t.rows[0].cells
for i, htxt in enumerate(["Misconception", "Guidance / feedback point"]):
    hdr[i].width = widths[i]
    p = hdr[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(htxt); r.font.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(hdr[i], GREEN_HEX)
for a, b in misodes:
    cells = t.add_row().cells
    cells[0].width = widths[0]; cells[1].width = widths[1]
    p = cells[0].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(a); r.font.size = Pt(8.8); r.font.bold = True
    p = cells[1].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(b); r.font.size = Pt(8.8)

# footer with page numbers
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Oak International School — MYP 5 Mathematics — Mid-Term Test 2 — Teacher's Mark Scheme    |    Page ")
r.font.size = Pt(7); r.font.color.rgb = GREY
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
run_el = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '14'); rpr.append(sz)
col = OxmlElement('w:color'); col.set(qn('w:val'), '6B7280'); rpr.append(col)
run_el.append(rpr)
t_el = OxmlElement('w:t'); t_el.text = '1'; run_el.append(t_el)
fld.append(run_el); fp._p.append(fld)
r = fp.add_run(" of "); r.font.size = Pt(7); r.font.color.rgb = GREY
fld2 = OxmlElement('w:fldSimple'); fld2.set(qn('w:instr'), 'NUMPAGES')
run_el2 = OxmlElement('w:r'); rpr2 = OxmlElement('w:rPr')
sz2 = OxmlElement('w:sz'); sz2.set(qn('w:val'), '14'); rpr2.append(sz2)
col2 = OxmlElement('w:color'); col2.set(qn('w:val'), '6B7280'); rpr2.append(col2)
run_el2.append(rpr2)
t_el2 = OxmlElement('w:t'); t_el2.text = '2'; run_el2.append(t_el2)
fld2.append(run_el2); fp._p.append(fld2)

doc.core_properties.title = "MYP 5 Mathematics — Mid-Term Test 2 — Teacher's Mark Scheme"
doc.core_properties.author = "Oak International School"
out = os.path.join(HERE, "MYP5_Mathematics_MidTerm_Test2_Mark_Scheme.docx")
doc.save(out)
print("Mark scheme 2 written:", out)
